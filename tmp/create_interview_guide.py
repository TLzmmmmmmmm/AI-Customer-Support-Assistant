from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\AI-Customer-Support-Assistant")
OUT = ROOT / "deliverables" / "Shengborun_AI_Project_Interview_Guide_ZH_EN.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.9)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)


def set_font(style, size, bold=False):
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr, value in (("ascii", "Aptos"), ("hAnsi", "Aptos"), ("eastAsia", "Microsoft YaHei")):
        fonts.set(qn(f"w:{attr}"), value)


styles = doc.styles
set_font(styles["Normal"], 9.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.18
for name, size, before, after in (
    ("Title", 21, 0, 13),
    ("Heading 1", 15, 15, 8),
    ("Heading 2", 11.5, 10, 5),
    ("Heading 3", 10, 8, 4),
):
    set_font(styles[name], size, True)
    fmt = styles[name].paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = True


def p(text="", *, label=None, style=None):
    para = doc.add_paragraph(style=style)
    if label:
        run = para.add_run(label)
        run.bold = True
    para.add_run(text)
    return para


def bi(zh, en, *, label_zh="中文  ", label_en="English  "):
    p(zh, label=label_zh)
    p(en, label=label_en)


def h1(zh, en):
    doc.add_heading(f"{zh}  {en}", level=1)


def h2(zh, en):
    doc.add_heading(f"{zh}  {en}", level=2)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        b.append(e)
    tcpr.append(b)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
        shade(t.rows[0].cells[i], "DDEBF7")
        t.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            borders(cell)
            if widths:
                cell.width = Cm(widths[i])
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(8.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def bullet(zh, en):
    p(zh, label="• 中文  ")
    p(en, label="  English  ")


title = doc.add_paragraph(style="Title")
title.add_run("盛博润 AI 客服项目面试手册\nShengborun AI Support Interview Guide")
p("项目流程、系统设计、实现证据与中英双语面试准备", label="用途  ")
p("Project flow, system design, implementation evidence, and bilingual interview preparation", label="Purpose  ")
bi(
    "这份手册把企业官网与 AI 客服后端作为一个集成项目来解释。面试时先说明解决了什么客户问题，再说明为何按请求类型选择检索、精确查询或工具，最后用限定范围的评测结果证明工程效果。",
    "This guide treats the company website and AI support backend as one integrated project. In an interview, start with the customer problem, explain why different request types take different retrieval or tool paths, and close with evaluation results and their limits.",
)

h1("一 项目概况", "1 Project profile")
table(
    ["项目维度 / Dimension", "核对结果 / Verified project record"],
    [
        ("范围 / Scope", "Astro/TypeScript 静态官网 + FastAPI/Python AI 客服 / static company site plus AI backend"),
        ("业务内容 / Content", "49 产品、4 类别、6 解决方案 / 49 products, 4 categories, 6 solutions"),
        ("知识快照 / Knowledge", "62 文档、74 知识块、74 向量，每条 1,024 维 / 62 documents, 74 chunks, 74 vectors, 1,024 dimensions"),
        ("核心技术 / Stack", "Astro, TypeScript, Zod, FastAPI, Pydantic, LangGraph, DashScope Embeddings, NumPy, DeepSeek"),
        ("部署记录 / Deployment record", "Ubuntu, Nginx, HTTPS; FastAPI/Uvicorn behind Nginx; backend systemd service"),
    ],
    [4.1, 12.4],
)
bi(
    "项目记录将 AI 后端工作列为 2026 年 8 至 9 月。仓库可证明实现与评测内容，不能单独证明劳动关系、自然客户流量、销售转化或客服人力节省。",
    "The project record dates backend work to August–September 2026. Repository evidence supports implementation and evaluation claims, but does not by itself prove employment status, organic customer traffic, revenue impact, or support labor savings.",
)

h1("二 从问题到系统", "2 Problem and evolution")
h2("业务问题", "Business problem")
bi(
    "访客可能询问具体型号参数、推荐适用产品、公司方案、联系方式或简单问候。这些请求需要不同的可靠性保证：已知型号应精确查询；不知道哪款产品合适才需要语义排序；联系方式应来自确定性数据源；无依据的问题应安全回退。",
    "Visitors ask for exact model specifications, suitable product recommendations, company solutions, contact details, or simple conversation. Each needs a different guarantee: known identities deserve exact lookup, discovery calls for semantic ranking, contact facts should come from deterministic data, and unsupported claims require safe fallback.",
)
h2("迭代流程", "Evolution")
table(
    ["阶段 / Stage", "问题与实现 / Problem and implementation"],
    [
        ("V0", "仅 LLM + 保守提示词；20 题脚本基线 2/20 正确，17 个可回答的公司问题被拒答 / LLM only; 2/20 scripted baseline, 17 answerable company questions refused"),
        ("V1", "规范化企业知识、结构化分块、本地 RAG，让公司事实有可检索证据 / normalized company knowledge, structural chunks, local RAG"),
        ("V2", "精确实体解析与三个只读工具，避免将型号和联系方式交给近似检索 / exact entity resolution and three read-only tools"),
        ("V3", "六路径混合路由 + LangGraph，确定性路径先行，复杂请求进入受约束 Agent / six-route hybrid router and LangGraph"),
        ("V4", "请求级遥测、分层评测、安全边界测试和性能发布门槛 / telemetry, layered evaluation, trust-boundary checks, performance gate"),
    ],
    [2.1, 14.4],
)

h1("三 架构与请求流程", "3 Architecture and request flow")
p("Astro 页面与聊天组件 → POST /api/chat-stream → 请求校验与准入控制 → HybridRouter → 检索 / 精确工具 / Agent / 直接回答 / 回退 → 答案清理与引用 → NDJSON 事件 → 前端展示", label="中文路径  ")
p("Astro UI → POST /api/chat-stream → validation and admission → HybridRouter → retrieval / exact tool / agent / direct / fallback → answer sanitation and citations → NDJSON events → UI", label="English flow  ")

h2("官网与聊天前端", "Website and chat frontend")
bi(
    "官网使用 Astro Content Collections 加载 JSON 产品和类别、Markdown 解决方案，并用 Zod Schema 校验字段。共享工具筛选已发布内容，动态路由在构建时生成静态页面；另有引用关系校验检查重复 ID、类别、slug 和产品特性。SEO 包含 canonical、Open Graph 与 sitemap。聊天组件处理历史消息、输入限制、错误和引用展示。",
    "The site loads product and category JSON plus solution Markdown through Astro Content Collections with Zod schemas. Shared helpers select published content and build static routes. A separate validator checks duplicate IDs, category relations, slugs, and feature references. SEO includes canonical URLs, Open Graph metadata, and a sitemap. The chat widget handles history, input limits, errors, and citations.",
)
bi(
    "前端按行解析 NDJSON 的 delta、citations、done 和 error 事件。当前后端先完成模型调用，再一次发送完整答案的 delta，因此传输协议支持事件，但线上路径不是逐 Token 输出。只有收到完整答案和 done，前端才把该轮加入会话历史。",
    "The frontend parses NDJSON delta, citations, done, and error events line by line. The current backend completes model generation before emitting one full-answer delta, so the event transport is not live token streaming. The UI commits a turn to conversation history only after receiving a complete answer and done.",
)

h2("知识构建与检索", "Knowledge build and retrieval")
bi(
    "官网权威内容经人工审核同步到后端 knowledge/source；构建脚本验证来源模型与关系，生成带稳定 ID、哈希和来源路径的文档与结构化分块，再调用 DashScope 生成向量。统一构建先在临时目录完成并验证快照，之后才替换线上产物；未变更块仅在 ID、哈希、模型及维度等配置一致时复用向量。这个流程不会自动抓取官网。",
    "Owner-reviewed website content is synchronized into backend knowledge/source. Builders validate schemas and relationships, create documents and structural chunks with stable IDs, hashes, and provenance, then generate DashScope embeddings. A unified build validates a staged snapshot before replacing live artifacts. An unchanged vector is reused only when identity, hash, model, dimensions, and related settings match. The pipeline does not scrape the website automatically.",
)
bi(
    "在线检索默认 Top-K 5。查询向量进入内存 NumPy 精确余弦索引；识别到实体时优先返回对应父文档结果，再用稠密检索补足。精确产品参数路径则直接解析 canonical 产品 ID 并调用 get_product_details，不依赖向量相似度。74 条向量的规模尚不支持宣称大规模检索能力。",
    "Online retrieval defaults to Top-K 5. The query embedding is searched against an in-memory NumPy exact cosine index; entity matches are prioritized before dense results fill remaining slots. Exact specification requests instead resolve a canonical product ID and call get_product_details without semantic ranking. A 74-vector index does not establish large-scale retrieval performance.",
)

h2("路由 编排与工具", "Routing orchestration and tools")
table(
    ["路径 / Route", "处理方式 / Execution"],
    [
        ("direct", "简单对话，直接生成 / simple conversation"),
        ("knowledge", "检索公司、方案和支持资料后生成 / grounded company, solution, support answer"),
        ("exact_product", "解析型号并精确获取产品资料 / canonical product lookup"),
        ("product_search", "按需求检索候选产品，必要时走 Agent / semantic discovery and agent when needed"),
        ("contact", "确定性联系方式工具 / deterministic contact tool"),
        ("fallback", "不支持或无法验证的请求安全回退 / safe fallback"),
    ],
    [3.0, 13.5],
)
bi(
    "HybridRouter 先用确定性线索和实体解析处理高置信请求，歧义请求才交给只能返回六个合法标签之一的模型分类器。LangGraph 负责 route、retrieve、generate、execute_tool 和 finalize 等节点之间的状态流转；只有需要推理或多步工具选择时才走 Agent 分支。",
    "HybridRouter uses deterministic cues and entity resolution for high-confidence cases. Ambiguous cases use an LLM classifier restricted to six valid labels. LangGraph moves state through route, retrieve, generate, execute_tool, and finalize nodes. The agent branch is reserved for requests that need reasoning or tool selection.",
)
bi(
    "三个注册工具是 search_products、get_product_details 与 get_contact_info，均只读。工具执行器检查允许列表、JSON 对象格式、Pydantic 参数和返回类型；同一请求最多处理三次工具调用，超预算批次在执行前整体拒绝。成功观察可按规范化参数复用。",
    "The registered read-only tools are search_products, get_product_details, and get_contact_info. The executor validates the allowlist, JSON object shape, Pydantic arguments, and result type. Each request has a three-call tool budget; an over-budget batch is rejected before any call runs. Successful observations can be reused by normalized arguments.",
)

h2("信任边界与运行保障", "Trust boundaries and operations")
bi(
    "HTTP 层验证角色交替、单条和会话长度；进程内限流与并发信号量限制进入模型路径的请求。客户端提交的历史消息、检索文本、工具观察和模型输出都不能被视为系统指令。最终回答清理不可信链接，只附加通过来源校验且实际被答案引用的资料；公开错误不包含内部异常细节。",
    "The HTTP layer validates role order and message/conversation sizes. Process-local rate limiting and a semaphore control admission. Client-supplied history, retrieved text, tool observations, and model output are not trusted instructions. Finalization sanitizes untrusted links and includes only validated sources actually cited by the answer. Public errors hide internal exception details.",
)
bi(
    "部署记录显示静态官网经 Nginx/HTTPS 提供服务，FastAPI/Uvicorn 运行在 Ubuntu 上并由 Nginx 代理。后端知识更新需先构建并验证完整快照，再重启服务。当前限流和并发状态仅在单进程内有效；多实例时需要共享基础设施。",
    "Deployment records describe the static site behind Nginx/HTTPS and FastAPI/Uvicorn on Ubuntu behind an Nginx proxy. Knowledge changes are built and validated as a complete snapshot before service restart. Rate and concurrency controls are process-local and would need shared infrastructure for multiple instances.",
)

h1("四 证据与结果的正确说法", "4 Evidence and claim boundaries")
table(
    ["指标 / Metric", "结果 / Result", "面试限定 / Interview qualifier"],
    [
        ("检索 Hit@5 / Retrieval", "Dev 23/23; Frozen 18/18; Holdout10 9/9", "历史内部题集，证明目标证据进入前五名；不是线上回答正确率 / historical internal sets, not online answer accuracy"),
        ("路由 / Routing", "24/24 final Dev", "24 个开发集用例，不是全域准确率 / 24 development cases"),
        ("工具动作 / Tool action", "raw 23/24; reviewed 24/24", "LY198 大小写差异经人工复核 / one case adjudicated for canonical ID case"),
        ("答案质量 / Answer quality", "47/48 manual Dev score", "有一次轻微无依据推断 / one partially unsupported inference"),
        ("代表性请求 / Representative workload", "132/132 success; P50 1.00 s; P95 3.156 s", "顺序、受控、生产路径风格；非自然流量或 SLO / sequential controlled workload, not organic traffic or SLO"),
        ("流式实验 / Streaming gate", "TTFT -42.4%; total latency +16.9%", "产品搜索路径；超过 15% 回退阈值，未发布 / product-search route, failed 15% guardrail, not shipped"),
        ("安全 / Security", "32 sampled runs; 0 reviewed breaches", "内部对抗评估，非渗透测试或安全认证 / internal red-team sample, not certification"),
        ("回归测试 / Regression", "669 backend in final audit; 28 frontend unit", "后端数字来自既有审计记录；前端 28 项与构建已在当前代码复核 / backend historical audit; frontend current check"),
    ],
    [3.0, 5.0, 8.5],
)
bi(
    "不要把 132 条顺序测试称为高并发压测，不要把 Hit@5 称为客户满意度，不要声称 64/64 浏览器 E2E 在最终审计中重新通过，也不要说系统已上线逐 Token 流式、模型微调、向量数据库或可写 CRM 工具。",
    "Do not call the 132 sequential requests a high-concurrency load test, translate Hit@5 into customer satisfaction, claim that all 64 browser E2E cases were rerun in the final audit, or claim shipped token streaming, model fine-tuning, a vector database, or write-capable CRM tools.",
)

doc.add_page_break()
h1("五 面试官问题与追问", "5 Interview questions and follow-ups")
p("每题同时给出中文与英文问法、可能追问和回答抓手。抓手是准备要点，不是背诵稿。", label="使用方式  ")
p("Each entry includes the main question, likely follow-ups, and evidence to use in an answer. These are prompts for preparation, not a script to memorize.", label="How to use  ")

questions = [
    ("项目范围 / Ownership", "你具体负责哪些部分？官网与 AI 后端是同一个系统吗？", "What did you personally own, and how do the website and AI backend relate?", "哪些工作是你独立完成的？哪些依赖公司资料审核？", "What did you build yourself, and what required company content review?", "两个仓库、前后端边界、资料人工审核同步；不扩大个人贡献。", "Name the two repositories, their boundary, and owner-reviewed content sync without overstating ownership."),
    ("项目范围 / Ownership", "为什么先做官网，再做 AI 客服？", "Why build the website before the AI assistant?", "网站内容怎样变成客服知识？为什么不直接爬站？", "How did site content become assistant knowledge, and why not scrape it?", "官网建立权威内容和稳定路由，后端用审核后的 knowledge/source 快照。", "The site provided authoritative content and stable routes; the backend used reviewed source snapshots."),
    ("产品与前端 / Product and frontend", "为什么选 Astro 静态站，而不是传统 SPA 或 SSR？", "Why Astro static generation instead of an SPA or SSR?", "新增产品会改几处代码？SEO 与构建校验怎么做？", "What changes when adding a product, and how do you protect SEO and content quality?", "内容展示为主；Content Collections、Schema、生成路由、静态交付。", "Content-heavy site; collections, schemas, generated routes, static delivery."),
    ("产品与前端 / Product and frontend", "聊天窗口如何处理流事件和失败消息？", "How does the chat widget handle stream events and failed turns?", "断流后是否写入历史？什么时候展示引用？", "Does an interrupted turn enter history? When are citations shown?", "按行解析 NDJSON；done 结束；失败时删除不完整回复；完成才入历史。", "Line-based NDJSON parsing; done terminates; failed partial response is removed; history commits on completion."),
    ("架构 / Architecture", "最初只有提示词的方案哪里失败了？", "What failed in the initial prompt-only version?", "2/20 是什么题集？为什么拒答比编造更好？", "What was the 2/20 dataset, and why were refusals preferable to fabrication?", "17 个可答公司问题因无权威知识被拒；补知识访问而非放松安全提示词。", "Seventeen answerable company questions were refused without knowledge access; add grounding rather than weaken safeguards."),
    ("架构 / Architecture", "为何需要六条路由，而不是所有请求都交给 Agent？", "Why use six routes instead of sending every request to an agent?", "混合意图和模糊请求如何处理？分类器输出非法标签怎么办？", "How do you handle mixed intent and invalid classifier labels?", "已知事实用确定性路径；歧义才用受约束分类；非法输出回退。", "Use deterministic paths for known cases; constrained LLM for ambiguity; invalid outputs fall back."),
    ("架构 / Architecture", "为什么精确型号查询不走普通向量搜索？", "Why should an exact model-number query bypass ordinary vector search?", "别名、大小写、相似型号怎么处理？", "How do aliases, case differences, and similar model IDs work?", "身份已知应解析 canonical ID；相似度用于发现未知候选。", "Resolve known identity to a canonical ID; rank semantically only for discovery."),
    ("架构 / Architecture", "LangGraph 在这里解决了什么问题？", "What did LangGraph solve in this system?", "你原来有直接工具调用循环吗？图中什么条件会回到工具节点？", "Did you first build a raw tool loop, and what causes another graph tool step?", "先验证原生工具语义，再用显式节点与边统一路径；不要说是多 Agent 系统。", "Validated native tool semantics first, then encoded paths as graph nodes and edges; this is not a multi-agent system."),
    ("知识工程 / Knowledge", "知识源如何规范化并保留可追溯性？", "How did you normalize knowledge while preserving provenance?", "网站字段变化、重复 ID、来源 URL 失效怎么办？", "What happens with site schema changes, duplicate IDs, or stale source URLs?", "来源 Schema、关系校验、稳定 ID、内容哈希、source_files/source_url。", "Discuss schemas, relationship checks, stable IDs, hashes, and source paths/URLs."),
    ("知识工程 / Knowledge", "为什么采用结构化分块？如何避免丢失产品上下文？", "Why use structural chunks, and how do you preserve product context?", "固定长度分块或更大块会怎样？", "What would fixed-size or larger chunks change?", "按业务文档和章节边界切分，保留父文档与来源；用检索题集验证。", "Split by business entity and section, retain parent and source, and evaluate retrieval."),
    ("知识工程 / Knowledge", "为什么用 NumPy 余弦索引而不用向量数据库？", "Why NumPy exact cosine search instead of a vector database?", "语料扩大一千倍时你的迁移信号是什么？", "What metrics would trigger a migration at much larger scale?", "74 条向量的运维复杂度不值得；观察内存、P95、更新频率和并发。", "At 74 vectors a service adds little; watch memory, P95, update cadence, and concurrency."),
    ("知识工程 / Knowledge", "知识更新如何避免半成品上线？", "How do you avoid activating a partial knowledge update?", "什么条件下旧向量可复用？外部 embedding 调用失败怎么办？", "When can a vector be reused, and what if embedding fails?", "临时目录构建、完整验证后替换；匹配 ID、hash、provider/model/dimensions。", "Build in staging, validate complete snapshot, then replace; reuse only with matching identity/hash/config."),
    ("工具与安全 / Tools and security", "三个工具的职责及边界是什么？", "What do the three tools do, and what are their boundaries?", "为什么只读？新增写入工单工具需要什么？", "Why read-only, and what would a ticket-writing tool require?", "产品搜索、精确详情、联系方式；写能力需鉴权、确认、幂等、审计。", "Product discovery, exact details, contact info; writes need auth, confirmation, idempotency, audit."),
    ("工具与安全 / Tools and security", "模型提出未知工具或错误参数时会怎样？", "What happens when the model proposes an unknown tool or malformed arguments?", "三次预算如何防止批量调用绕过？是否会先执行一部分？", "How is the three-call budget enforced for batches? Can part of a rejected batch execute?", "允许列表、Pydantic 严格模型、超预算批次原子拒绝。", "Allowlist, strict Pydantic schemas, atomic pre-execution batch rejection."),
    ("工具与安全 / Tools and security", "怎样防止检索文本或工具返回值中的提示注入？", "How do you handle prompt injection in retrieved text or tool observations?", "哪些安全性质靠代码，哪些仍依赖模型？", "Which properties are code-enforced, and which still depend on model behavior?", "数据与控制分离；工具授权、参数、预算由代码执行；语义回答仍需评测。", "Separate data from control; enforce tool auth/schema/budget in code; evaluate model-dependent answer properties."),
    ("工具与安全 / Tools and security", "引用链接能否由模型直接生成？", "Can the model invent citation URLs?", "引用如何去重？如果答案只提到一条来源怎么办？", "How are citations deduplicated, and what if an answer cites only one source?", "来源对象校验、规范 URL 去重、稳定 source ID、只呈现实际引用的来源。", "Validate sources, normalize URLs for dedupe, use stable IDs, display only actually cited sources."),
    ("工具与安全 / Tools and security", "客户端传来的历史消息可信么？", "Can you trust client-supplied conversation history?", "如果伪造先前 assistant 授权怎么办？", "What if a user forges a prior assistant authorization?", "历史可校验格式但非加密认证；不能覆盖系统政策或授权边界。", "History is shape-validated but not authenticated; it cannot override policy or tool authority."),
    ("可靠性 / Reliability", "超时、限流、并发和模型故障如何处理？", "How do timeout, rate limits, concurrency, and provider failures behave?", "多实例部署时当前限流还有效吗？", "Would current admission controls work across multiple instances?", "请求校验先于编排；进程内限制；公共错误脱敏；多实例需共享状态。", "Validate before orchestration; process-local controls; redact errors; share state for multi-instance."),
    ("评测 / Evaluation", "Hit@5 到底衡量什么？", "What does Hit@5 actually measure?", "为什么它不能证明最终回答正确？Dev 与 Holdout 如何隔离？", "Why does it not establish final-answer correctness, and how were sets separated?", "目标证据进入前五名；路由、工具和人工答案质量另测。", "It measures evidence presence in the top five; routing, actions, and answer quality are separate."),
    ("评测 / Evaluation", "24/24 工具动作正确，有没有特殊判定？", "Was the 24/24 tool-action result purely automatic?", "原始 23/24 为什么变成 24/24？如何避免为了分数改金标？", "Why did raw 23/24 become reviewed 24/24, and how did you avoid moving the goalposts?", "ly198 与 LY198 是同一 canonical 实体；报告保留原始分数与人工复核记录。", "One case differed only by canonical-ID letter case; preserve raw score and adjudication trail."),
    ("评测 / Evaluation", "47/48 答案质量失分说明什么？", "What does the 47/48 answer-quality score reveal?", "正确工具结果仍会产生什么风险？", "What risk remains even when route and tool are correct?", "酒店场景出现证据未支持的轻微推断；生成质量独立于动作正确性。", "A hotel recommendation added a lightly unsupported inference; generation quality is separate from action correctness."),
    ("性能 / Performance", "132/132 成功和 P95 3.156 秒能证明什么？", "What do 132/132 success and 3.156-second P95 establish?", "是线上 SLO、并发压测还是顺序合成工作负载？", "Was this an online SLO, concurrency load test, or sequential synthetic workload?", "受控生产路径风格顺序请求；提供基线，不代表自然流量。", "A controlled sequential production-path baseline, not organic traffic or load-test capacity."),
    ("性能 / Performance", "你怎么判断瓶颈在模型而非检索？", "How did you determine the model, not retrieval, dominated latency?", "产品搜索路径的模型与工具中位数分别多少？", "What were median model and tool times for product search?", "产品搜索总 P50 2235ms，模型约 1984ms，工具/检索约 157ms。", "Product-search P50 2,235 ms; model about 1,984 ms; tool/retrieval about 157 ms."),
    ("性能 / Performance", "为什么流式输出变快了却没有发布？", "Why was token streaming not shipped despite a faster first token?", "实验设计、样本数、发布门槛和总延迟怎么定义？", "What were the experiment design, sample size, gate, and total-latency definition?", "36 个受控请求；产品搜索 TTFT -42.4%，总延迟 +16.9%，超过 15% 门槛。", "36 controlled requests; product-search TTFT -42.4%, total latency +16.9%, above 15% guardrail."),
    ("安全 / Security", "32 次红队运行零突破意味着什么？", "What does zero breaches in 32 red-team runs mean?", "是否等于安全认证？残余风险是什么？", "Is this a security certification? What remains risky?", "仅在 12 类内部攻击样本中未见边界突破；模型语义行为与新数据源需持续复查。", "No observed breach in 12 internal attack families; model semantics and new sources still need review."),
    ("运营与扩展 / Operations and scaling", "如何发布并回滚官网和知识快照？", "How do you deploy and roll back the site and knowledge snapshot?", "构建成功是否足够？何时重启后端？", "Is a successful build sufficient, and when do you restart the backend?", "官网静态构建与 Nginx 切换；后端先验证完整知识快照再重启。", "Static build and Nginx release; validate full knowledge snapshot before backend restart."),
    ("运营与扩展 / Operations and scaling", "如果产品从 49 款变成数万款，你先改什么？", "What would change first if the catalog grew from 49 to tens of thousands?", "检索基础设施、内容同步、缓存和测量如何演进？", "How would retrieval, content sync, caching, and measurement evolve?", "基于实际 P95、索引内存、更新频率与召回质量判断；不要先假设一定要向量库。", "Use measured P95, memory, update cadence, and recall quality before choosing new infrastructure."),
]

last_category = None
for idx, item in enumerate(questions, 1):
    category, qzh, qen, fzh, fen, azh, aen = item
    if category != last_category:
        doc.add_heading(category, level=2)
        last_category = category
    doc.add_heading(f"Q{idx:02d}  {qzh}", level=3)
    p(qen, label="Question  ")
    p(fzh, label="追问  ")
    p(fen, label="Follow-up  ")
    p(azh, label="作答抓手  ")
    p(aen, label="Answer anchor  ")

doc.add_page_break()
h1("六 一至两分钟项目介绍", "6 One to two minute project pitch")
h2("中文版", "Chinese version")
p(
    "我在盛博润项目中负责企业官网和 AI 客服系统的开发与集成。官网使用 Astro 和 TypeScript，把 49 款产品、4 个产品分类和 6 个行业解决方案整理成可校验的结构化内容，并生成静态页面。\n\n做 AI 客服时，我先验证了仅靠提示词的基线：模型没有公司权威资料，很多本来能回答的问题只能拒答。因此我构建了从审核后的来源数据到文档、知识块和向量的流程，目前有 62 份文档、74 个知识块。对一般公司知识问题使用 RAG；对明确型号和联系方式使用精确解析与只读工具。之后我设计了确定性优先的六路径路由，并用 LangGraph 编排需要工具或检索的流程。\n\n我把检索、路由、工具动作和最终答案分开评测。内部检索题集的 Hit@5 全部命中；最终 Dev 集中路由 24/24 正确，工具行为 24/24 经复核符合预期。在 132 条受控代表性请求中全部成功，P50 为 1 秒、P95 为 3.156 秒。这些是内部测试，不是自然客户流量。项目让我形成的主要方法是：身份明确时用确定性查询，相关性未知时才用语义检索，并用完整指标决定是否发布优化。"
)
h2("English version", "English version")
p(
    "I worked on both the company website and the AI customer-support system for Shengborun. On the frontend, I used Astro and TypeScript to organize 49 products across four categories and six solution pages as validated content, then generated a static site.\n\nFor the assistant, I first tested a prompt-only baseline. The model had no authoritative company knowledge, so it refused many questions that should have been answerable. I built a reviewed knowledge pipeline that now contains 62 normalized documents and 74 chunks. General company questions use RAG, while known model numbers and contact details use exact resolution and read-only tools. I then added a deterministic-first router with six routes and used LangGraph to orchestrate retrieval and tool paths when needed.\n\nI evaluated retrieval, routing, tool actions, and final answers separately. The internal retrieval sets reached complete Hit@5 coverage, and the final development set had 24 out of 24 correct routes and 24 out of 24 tool actions after a documented adjudication. A controlled 132-request workload completed successfully, with one-second median and 3.156-second P95 latency. Those numbers describe internal tests, not live customer traffic. My central design lesson was to use exact lookup when identity is known, semantic ranking when relevance is uncertain, and full performance guardrails before shipping an optimization."
)

h1("七 可复用面试故事", "7 Reusable interview stories")
p("每个故事按情境、任务、行动、结果组织。讲述时可根据问题压缩至 45 至 90 秒。", label="使用方式  ")
p("Each story follows situation, task, action, and result. Compress it to 45–90 seconds based on the question.", label="How to use  ")

stories = [
    (
        "故事一 从安全拒答到有依据的回答", "Story 1 From safe refusals to grounded answers",
        "情境：最初的客服只有模型和保守提示词。20 题脚本基线仅 2 题正确，17 个本可由公司资料回答的问题被拒答。任务：改善可回答性，同时不允许模型编造产品事实。行动：先规范化官网资料，再按业务实体和章节分块，生成向量并建立来源引用；用独立检索题集检查证据能否进入 Top 5。结果：当前快照包含 62 个文档和 74 个知识块，历史内部评测 Hit@5 为 Dev 23/23、Frozen 18/18、Holdout10 9/9。关键判断：问题是缺少权威知识访问，而不是提示词不够大胆。",
        "Situation: The first assistant used only an LLM and a conservative prompt. It got 2 of 20 scripted baseline cases correct and refused 17 answerable company questions. Task: Improve answerability without fabricating company facts. Action: I normalized reviewed website content, created business-entity and section-based chunks, embedded them, retained source provenance, and measured whether expected evidence appeared in the top five. Result: The current snapshot has 62 documents and 74 chunks; historical internal Hit@5 was 23/23 on Dev, 18/18 on Frozen, and 9/9 on Holdout10. The core diagnosis was missing authoritative knowledge access, not an overly cautious prompt."
    ),
    (
        "故事二 用内容模型交付可维护的官网", "Story 2 Making the website maintainable through content models",
        "情境：官网要展示大量专业通信产品与长篇方案，逐页手工维护容易产生重复和错链。任务：在内容规模增长时仍保持页面一致和可更新。行动：用 Astro Content Collections、Zod Schema 与共享路由工具组织产品、分类和方案；增加跨内容引用校验，构建时拦截重复 ID、无效分类及产品特性。结果：网站覆盖 49 款产品、4 个分类和 6 个方案；当前代码的生产构建与 28 项单测通过。关键判断：内容变更应主要是数据维护，而不是复制页面代码。",
        "Situation: The site needed a large technical catalog and long solution pages, making hand-maintained pages prone to duplication and broken links. Task: Keep pages consistent and easy to update as content grew. Action: I used Astro Content Collections, Zod schemas, shared routing helpers, and cross-content validation for IDs, categories, and product features. Result: The site covers 49 products in four categories and six solutions; the current production build and 28 unit tests pass. The design choice was to make most catalog updates data work rather than duplicated page code."
    ),
    (
        "故事三 修复多轮对话里的指代错误", "Story 3 Fixing a contextual follow-up failure",
        "情境：开发集里，用户先问 LY198 功率，再问“那它支持什么频段？”。最初系统看不到第二句里的型号，因此回退。任务：让多轮问题沿用正确产品身份，同时避免猜错。行动：路由时识别指代词，并从此前用户消息中解析产品 ID；仅有一个明确历史产品时走精确产品路径，多个候选时保持 Agent 分支处理歧义。结果：相关 Dev 用例修复，最终路由评测为 24/24。关键判断：会话上下文能帮助解析身份，但不能绕过实体唯一性检查。",
        "Situation: In a development case, the user asked for LY198's power and then, 'What frequency band does it support?' The second turn omitted the model ID and initially fell back. Task: Preserve the right product identity without guessing. Action: I detected contextual references and resolved product IDs from earlier user messages. A single historical product takes the exact-product path; multiple candidates stay on an agent path for ambiguity. Result: The case was repaired and final Dev routing scored 24/24. Conversation context helps resolve identity, but only when ambiguity is handled explicitly."
    ),
    (
        "故事四 不发布看起来更快的优化", "Story 4 Declining to ship an apparently faster optimization",
        "情境：性能分析发现产品搜索最慢，模型耗时远高于本地工具检索。我尝试逐 Token 流式输出以缩短首字等待。任务：改善体验且不显著牺牲总完成时间。行动：预先设定首 Token 和总延迟门槛，在 36 条受控请求中比较基线与流式方案。结果：产品搜索首 Token 中位时间改善 42.4%，但总延迟回退 16.9%，超过 15% 发布上限；因此保留现有缓冲式回答。关键判断：优化要看完整门槛，不用单个漂亮指标掩盖退步。",
        "Situation: Product search was the slowest route, with model time dominating local retrieval time. I tried token streaming to reduce time to first token. Task: Improve perceived latency without materially harming completion time. Action: I set first-token and total-latency gates in advance and compared buffered and streaming paths across 36 controlled requests. Result: Product-search median time to first token improved 42.4%, but total latency regressed 16.9%, above the 15% release ceiling. I retained the buffered production path. One attractive metric did not override the full release criteria."
    ),
    (
        "故事五 把 Agent 安全放在代码边界上", "Story 5 Enforcing agent safety at code boundaries",
        "情境：客服会读取用户历史、检索内容和工具结果，这些都可能携带恶意指令。任务：防止文本内容转变成执行权限。行动：在 HTTP 层验证输入；在 ToolExecutor 里固定工具允许列表、参数 Schema 与三次调用预算；只暴露只读工具，并对公开错误及引用做清理。随后分别测试确定性边界与 12 类语义攻击。结果：32 次内部对抗运行中，经复核未见受保护边界突破。关键判断：这个结果不是安全认证，未来接入外部资料或可写工具必须重新审计。",
        "Situation: The assistant reads user history, retrieved text, and tool results, any of which may carry malicious instructions. Task: Prevent that text from becoming execution authority. Action: I validated HTTP input, enforced a tool allowlist, typed arguments, and a three-call budget in ToolExecutor, kept business tools read-only, and sanitized public errors and citations. I then tested deterministic boundaries and 12 semantic attack families. Result: No reviewed protected-boundary breach appeared in 32 internal adversarial runs. This is not a security certification; external sources or write-capable tools would require a new audit."
    ),
]

for zh_title, en_title, zh_body, en_body in stories:
    h2(zh_title, en_title)
    p(zh_body, label="中文  ")
    p(en_body, label="English  ")

doc.add_page_break()
h1("八 核对资料与代码索引", "8 Evidence and code index")
bi(
    "以下路径用于面试前回看实现。相对路径分别以 D:\\Shengborun 和 D:\\AI-Customer-Support-Assistant 为根。生成产物、第三方依赖和密钥文件不作为设计证据。",
    "Use these paths to revisit implementation before an interview. Relative paths are rooted at D:\\Shengborun and D:\\AI-Customer-Support-Assistant respectively. Generated output, third-party dependencies, and secret files are not design evidence.",
)
table(
    ["主题 / Topic", "前端仓库 / Frontend", "后端仓库 / Backend"],
    [
        ("项目与技术 / Overview", "README.zh-CN.md; package.json; astro.config.mjs", "README.zh-CN.md; CASE_STUDY.md; requirements.txt"),
        ("内容与构建 / Content", "src/content.config.ts; src/lib/content-rules.ts; scripts/validate-content.mjs; src/pages/[category]/", "knowledge_pipeline/core.py; chunking.py; retrieval/records.py; scripts/build_knowledge.py"),
        ("在线请求 / Runtime", "src/components/chat/ChatWidget.astro; chat-transport.ts", "main.py; routes/chat.py; routing/router.py; agent_graph/graph.py; agent_graph/nodes.py"),
        ("检索与工具 / Retrieval and tools", "src/content/products/; src/content/solutions/", "knowledge_pipeline/retrieval/retriever.py; index.py; support_tools/service.py; agent/executor.py"),
        ("安全与引用 / Safety and citations", "tests/e2e/chat-widget.spec.ts; tests/unit/chat-transport.test.ts", "models.py; rate_limit.py; citation/core.py; tests/security/"),
        ("评测与部署 / Evaluation and deployment", ".github/workflows/ci.yml; DEPLOYMENT.md", "docs/week-3-day-6-agent-evaluation-final-report.md; docs/week-4-day-2-performance-report.md; docs/week-4-day-3-streaming-closeout.md; docs/security/week4-security-red-team-closeout.md; DEPLOYMENT.md"),
    ],
    [3.0, 6.6, 6.9],
)
bi(
    "面试时可以把任何一个数字指回其题集、报告或代码边界；对于未经测量的转化率、客服节省、生产在线准确率和高并发能力，直接说明项目未测量。",
    "In interviews, tie each number to its dataset, report, or code boundary. For conversion, support labor savings, live accuracy, or high-concurrency capacity, state plainly that this project did not measure them.",
)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Shengborun AI Support  |  Interview Guide  |  中英双语")
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "盛博润 AI 客服项目面试手册 Shengborun AI Support Interview Guide"
doc.core_properties.subject = "Bilingual project architecture and interview preparation"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
print(f"questions={len(questions)} stories={len(stories)}")
